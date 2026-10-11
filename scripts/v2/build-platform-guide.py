#!/usr/bin/env python3
"""Build the reference for the learning platform actually installed on 2026-10-10."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

root = Path(__file__).resolve().parents[2]
document = Document()
section = document.sections[0]
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.7)
normal = document.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(6)
for name in ['Heading 1', 'Heading 2']:
    document.styles[name].font.color.rgb = RGBColor.from_string('9E1527')

def para(text):
    document.add_paragraph(text)

def heading(text):
    document.add_heading(text, level=1)

def code(text):
    for line in text.splitlines():
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(8)

def table(headers, rows):
    t = document.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = text

document.add_heading('StreamFlix v2', level=0)
document.add_heading('Kind Cluster and Observability Setup', level=1)
para('Reference date: 10 October 2026. This guide describes the local platform installed and verified during the cluster reset. It separates running platform services from application features that still await CI deployment.')

heading('1. What was completed')
para('The learning cluster named streamflix-v2 was created with one control-plane and two workers. Its kubectl context is kind-streamflix-v2. All three nodes were Ready at completion.')
para('Ten Helm releases were deployed. Local PKI, public CA distribution and Grafana HTTPS were configured. Generated credentials were submitted directly to Kubernetes Secrets without printing their values. Application containers, application datastores and application Helm releases were not deployed; they remain the next GitHub Actions phase. Argo CD is a later phase.')
table(['Item', 'Current value'], [
    ['Repository', '/mnt/c/Users/sandeep.reddy01/Downloads/k8s-project'],
    ['Cluster / context', 'streamflix-v2 / kind-streamflix-v2'],
    ['kind / Kubernetes', 'kind 0.30.0 / Kubernetes 1.34.0'],
    ['Nodes', 'streamflix-v2-control-plane; streamflix-v2-worker; streamflix-v2-worker2'],
    ['Application namespace', 'streamflix-v2; prepared with Secrets, public CA and restricted Pod Security'],
    ['Platform namespaces', 'kube-system, ingress-system, cert-manager, observability, observability-v2, local-path-storage'],
    ['Grafana browser URL', 'https://grafana.streamflix-v2.local:9443'],
    ['Future application URL', 'https://streamflix-v2.local:9443; application not deployed yet'],
])

heading('2. Installed components and pinned charts')
table(['Release / namespace', 'Chart version', 'Purpose / Kubernetes resources'], [
    ['cilium / kube-system', '1.20.2', 'Pod networking and policy enforcement; agent DaemonSet, operator Deployment, RBAC and Cilium CRDs. kube-proxy retained.'],
    ['traefik / ingress-system', '41.7.1', 'Ingress controller; Deployment, ClusterIP Service, IngressClass and RBAC. Routes application and Grafana ingress.'],
    ['metrics-server / kube-system', '3.14.0', 'Deployment, Service and APIService; resource metrics for kubectl top and future HPAs.'],
    ['cert-manager / cert-manager', 'v1.20.4', 'Controller, webhook and cainjector Deployments; certificate/issuer CRDs and RBAC.'],
    ['trust-manager / cert-manager', 'v0.25.0', 'Public CA distribution; Deployment, webhook and Bundle CRD. Source trust namespace is cert-manager.'],
    ['kube-prometheus-stack / observability', '92.3.0', 'Prometheus Operator and kube-state-metrics Deployments, node-exporter DaemonSet, Grafana Deployment, Prometheus and Alertmanager StatefulSets, PVCs and monitoring CRDs.'],
    ['tempo / observability', '1.24.4', 'Trace backend; StatefulSet, PVC and Service. Tempo image 2.9.0. This legacy chart reports deprecation.'],
    ['loki-v2 / observability-v2', '18.3.0', 'Log backend; monolithic StatefulSet/PVC plus gateway Deployment/Service. Values pin Loki image 3.7.0; chart metadata may report a different appVersion.'],
    ['alloy-v2 / observability-v2', '1.13.1', 'API-based pod log collection; one Deployment, ConfigMap, ServiceAccount and discovery/pods-log RBAC. Alloy image v1.20.1.'],
    ['vpa / kube-system', '5.1.0', 'VPA CRDs and recommender Deployment only. Updater and admission controller disabled; application VPA mode is Off.'],
])
para('The chart versions are pinned in scripts/v2/install-platform.sh. Runtime appVersion labels do not necessarily equal an image explicitly overridden in values. Helm chart pinning does not mean every platform image is pinned by digest.')

heading('3. Log collection: Alloy, not Fluent Bit')
para('The retired cluster used Fluent Bit and Promtail. The replacement uses Grafana Alloy. The Fluent Bit/log-shipper chart remains in the repository as a historical artifact but is not installed in this new cluster. Fluent Bit can also send logs to Loki; this switch is an implementation choice, not a requirement imposed by Loki.')
code('Pod stdout/stderr\n  -> kubelet / Kubernetes pod-log API\n  -> Alloy discovery and loki.source.kubernetes\n  -> Loki HTTP push through loki-v2-gateway\n  -> Grafana Loki datasource and Explore')
para('Alloy discovers Pods across namespaces, labels namespace/pod/container/service, and reads logs using its ServiceAccount permissions. One API collector can collect across nodes without mounting /var/log or running a collector DaemonSet on every node. Its configuration is gitops/platform/observability-v2/alloy-values.yaml. The write endpoint is derived from Kubernetes DNS: http://loki-v2-gateway.observability-v2.svc.cluster.local/loki/api/v1/push.')
para('Collection was verified for cert-manager, ingress-system, kube-system, local-path-storage, observability and observability-v2. The application namespace has no application Pods yet, so application log coverage cannot be verified until CI deploys them.')

heading('4. Metrics, traces and alerts')
table(['Signal', 'Flow and current status'], [
    ['Resource metrics', 'kubelet -> metrics-server -> metrics.k8s.io -> kubectl top / future HPA. This is separate from Prometheus historical metrics.'],
    ['Monitoring metrics', 'node-exporter, kube-state-metrics, kubelet/cAdvisor and control-plane endpoints -> Prometheus -> Grafana. All 28 active targets were healthy at completion.'],
    ['Application metrics', 'Prepared /metrics endpoints -> ServiceMonitors -> Prometheus. Application ServiceMonitors and application PrometheusRules are installed with the future application release.'],
    ['Traces', 'Prepared Node/Python OpenTelemetry -> OTLP HTTP port 4318 -> Tempo -> Grafana. A synthetic span was accepted and retrieved; cross-service application traces still await deployment.'],
    ['Alerts', 'Prometheus evaluates rules -> Alertmanager groups/routes alerts. Platform rules are installed; app rules cover 5xx error rate, restarts and no available replicas. No email/Slack receiver or external notification delivery has been configured/tested.'],
])
para('Tempo has two relevant HTTP ports: 4318 for OTLP ingestion and 3200 for Grafana queries/readiness. Grafana uses http://tempo.observability.svc.cluster.local:3200. Application OTLP configuration uses http://tempo.observability.svc.cluster.local:4318. Confusing these ports breaks the datasource even when ingestion works.')
para('The pinned legacy Tempo chart runs successfully but reports deprecation. A move to the maintained community chart requires a validated configuration, image and storage upgrade; it was not silently applied during this setup.')

heading('5. Files to keep and how to reproduce the setup')
table(['File', 'Purpose'], [
    ['infra/kind/kind-v2.yaml', 'Three-node cluster; pinned node image digest; default CNI disabled; local metrics bindings.'],
    ['scripts/v2/install-platform.sh', 'Fetches pinned charts and installs all ten platform releases into the explicit kind context.'],
    ['gitops/platform/cilium-learning.yaml', 'Cilium Kubernetes IPAM; kube-proxy replacement disabled; one operator.'],
    ['gitops/platform/traefik-learning.yaml', 'Control-plane placement, ClusterIP Service and JSON access logs; no CNI hostPort dependency.'],
    ['gitops/platform/prometheus-learning.yaml', 'Persistent monitoring, generated Grafana Secret reference, TLS ingress and Tempo datasource.'],
    ['gitops/platform/observability-v2/*', 'Loki, Alloy and labeled Grafana Loki datasource ConfigMap.'],
    ['gitops/platform/tempo-learning.yaml', 'Explicit OTLP receiver bindings and persistent learning trace storage.'],
    ['scripts/v2/prepare-learning.sh', 'Namespace security labels, credential bootstrap, root CA and public trust synchronization.'],
    ['scripts/v2/bootstrap-grafana.py', 'Random Grafana admin credential stored directly in Kubernetes; existing Secret preserved.'],
    ['scripts/v2/expose-learning-metrics.py', 'Correction for metrics bindings on the initially created local cluster only.'],
    ['scripts/v2/check-network-policy.py', 'Disposable cross-node allow/deny/allow test with namespace cleanup.'],
])
para('Prerequisites are a running Docker Desktop daemon, kind, kubectl, Helm, Python with PyYAML, OpenSSL and access to image/chart registries. Run from the repository root. The cluster already exists; do not run kind create cluster again against the current name. Use the following sequence only when intentionally rebuilding an absent cluster.')
code('cd /mnt/c/Users/sandeep.reddy01/Downloads/k8s-project\nkind create cluster --config infra/kind/kind-v2.yaml\nkubectl config use-context kind-streamflix-v2\nbash scripts/v2/install-platform.sh\nbash scripts/v2/prepare-learning.sh\nkubectl label namespace observability streamflix.io/trust=true --overwrite\nkubectl wait certificate/streamflix-grafana-tls -n observability --for=condition=Ready --timeout=180s')
para('The installer waits for workloads and stops on errors. Grafana ingress is installed before the learning ClusterIssuer; its certificate becomes ready after prepare-learning.sh creates that issuer. Helm monitoring readiness does not prove the ingress certificate is ready, so wait for the Certificate explicitly.')
para('The initial cluster had loopback-only controller/scheduler/etcd/kube-proxy metrics listeners. They were corrected in the running cluster using the local metrics script. The saved kind configuration now includes those settings for future rebuilds. Metrics are exposed inside the Docker node network; no metrics ports are published to Windows. This configuration is for the local lab, not an internet-facing production control plane.')
code('# Existing local cluster only, if control-plane metrics remain unreachable:\npython3 scripts/v2/expose-learning-metrics.py')
para('The initial runtime cluster also retains unused Docker bindings for 8080/8443 from its creation. The final saved kind file omits them. Access uses a port-forward on 9443 in both cases, which was verified successfully.')

heading('6. Grafana access and public CA trust')
para('Add the following hostnames to local Windows hosts/DNS, normally C:\\Windows\\System32\\drivers\\etc\\hosts. Editing that file may require Windows administrator privileges. This is deployment DNS configuration; the application uses relative API paths and does not ask end users for service URLs.')
code('127.0.0.1 grafana.streamflix-v2.local streamflix-v2.local')
para('Export only the public root certificate. Copy it from WSL to Windows if importing it into the Windows/browser trusted-root store. Do not export or distribute streamflix-root-ca private-key material.')
code("kubectl get configmap streamflix-db-ca -n streamflix-v2 -o jsonpath='{.data.ca\\.crt}' > /tmp/streamflix-learning-ca.crt\nkubectl port-forward -n ingress-system svc/traefik 9443:443")
para('Keep the port-forward terminal running, then open https://grafana.streamflix-v2.local:9443. Trust the public learning CA in the browser/OS store used by the browser. The certificate subject must match the hostname. The temporary port-forwards used during validation were stopped after checks; no persistent forwarding process was left running.')
para('Grafana username is admin. Retrieve its generated password locally; do not commit or share the value. The custom credential Secret is grafana-admin, not the default Secret name printed in generic Helm notes.')
code("kubectl get secret grafana-admin -n observability -o jsonpath='{.data.admin-password}' | base64 -d")
para('Grafana has Prometheus, Tempo and Loki v2 datasources. Open Explore, select Loki v2 and use the queries below. Select a time window containing emitted logs; an idle Pod may have no new log lines in a short window.')
code('{namespace="kube-system"}\n{namespace="observability"}\n# After application deployment:\n{namespace="streamflix-v2", pod="kafka-0"}\n{namespace="streamflix-v2", service="catalog-service"}')

heading('7. Verification commands and recorded results')
code('kind get clusters\nkubectl config current-context\nkubectl get nodes\nkubectl top nodes\nhelm list -A\nkubectl get pods -A\nkubectl get certificates -A\nkubectl get clusterissuer streamflix-ca\nkubectl get bundle streamflix-db-ca\npython3 scripts/v2/check-network-policy.py')
table(['Check', 'Recorded result on 2026-10-10'], [
    ['Cluster identity', 'Only streamflix-v2 remained; old devops nodes were deleted.'],
    ['Nodes / containers', 'Three nodes Ready; all running platform containers Ready.'],
    ['Helm', 'All ten platform releases deployed.'],
    ['Resource metrics', 'kubectl top nodes returned CPU and memory metrics.'],
    ['Prometheus', '28 of 28 active scrape targets healthy after local metrics corrections.'],
    ['NetworkPolicy', 'Cross-node access initially worked, was blocked by a policy, then restored by explicit allow; test namespace deleted.'],
    ['Certificates / trust', 'Learning root and Grafana Certificate Ready; ClusterIssuer Ready; Bundle Synced.'],
    ['Grafana TLS', 'HTTPS /api/health succeeded with public CA and hostname verification; database reported ok.'],
    ['Logs', 'Loki returned labels and logs for all six active platform namespaces.'],
    ['Traces', 'Synthetic OTLP span was accepted on 4318 and retrieved through Tempo query API on 3200.'],
    ['Alertmanager', 'Readiness endpoint returned OK; external delivery was not configured.'],
])
para('To inspect Prometheus targets without opening a browser, query the service through the authenticated Kubernetes API proxy. This command reports the current state, which may differ from the historical snapshot during rollouts.')
code("kubectl get --raw '/api/v1/namespaces/observability/services/http:kube-prometheus-stack-prometheus:9090/proxy/api/v1/targets'\nkubectl get --raw '/api/v1/namespaces/observability-v2/services/http:loki-v2-gateway:80/proxy/loki/api/v1/labels'\nkubectl get --raw '/api/v1/namespaces/observability/services/http:tempo:3200/proxy/ready'")

heading('8. Troubleshooting reference')
table(['Symptom', 'Check / action'], [
    ['Nodes NotReady immediately after creation', 'Expected before Cilium starts. Check Cilium DaemonSet/operator Pods and events.'],
    ['ImagePullBackOff / DNS error', 'Describe the Pod and inspect registry hostname resolution inside its kind node. Repair Docker/WSL DNS; recreating an application Pod alone does not repair node DNS.'],
    ['kubectl top: Metrics API unavailable', 'Check metrics-server Deployment, logs and v1beta1.metrics.k8s.io APIService. The insecure kubelet TLS flag is local-kind-only.'],
    ['Controller or kube-proxy Prometheus targets down', 'Inspect target errors and listener bindings. Use the saved kind metrics patches on rebuild, or the guarded local correction script for this cluster.'],
    ['Grafana hostname not found', 'Check Windows/browser hosts/DNS and whether the WSL loopback port-forward is reachable.'],
    ['TLS warning', 'Check Certificate Ready, hostname, public CA trust and whether the port-forward targets Traefik TLS port 443.'],
    ['Missing Pod logs', 'Check the time window, actual kubectl logs output, Alloy discovery/RBAC and Loki datasource labels. Pods are listed from log streams, not from a complete live Pod inventory.'],
    ['Log-volume parse error', 'Verify the selected datasource points to the fresh Loki instance and inspect its actual image/query. The retired Loki datasource is not part of this setup.'],
    ['Tempo datasource failure', 'Use port 3200 for queries; port 4318 is ingestion. Check receiver binding, export endpoint and network policy separately.'],
    ['Alert shows firing but no message arrives', 'Alertmanager readiness is not delivery. Configure a receiver, route and credentials, then fire a controlled alert and verify delivery/recovery.'],
])

heading('9. Remaining work and operating limits')
para('Next: run GitHub Actions checks/builds, publish and promote image digests, configure the dedicated deployment runner and deploy the application release. Then verify real browser workflows, application metrics, Kafka/Postgres log streams, a gateway-to-service-to-datastore trace and application alert firing. Argo CD comes after direct CI deployment is validated.')
para('This is a learning platform. Loki filesystem storage, single-instance Tempo and node-local PVCs are not production HA or offsite recovery. Kubernetes Secrets are not encrypted merely because values are base64 encoded. Production still needs resilient data services, secret-manager/workload identity, encryption at rest, protected backups, restore drills, scoped network/RBAC policies, mesh/datastore TLS validation, and operating SLOs. The installed chart provides no external notification receiver and no production Istio or External Secrets controller was installed in this learning phase.')
para('Retain this setup guide together with docs/v2/DEPLOYMENT.md, docs/v2/VALIDATION.md and StreamFlix-v2-Architecture-and-Deployment.docx. The architecture document covers each microservice; this document covers the running cluster and platform setup.')

section.footer.paragraphs[0].text = 'StreamFlix v2 | Local platform setup | Verified 10 October 2026'
output = root / 'docs/v2/Kind-Cluster-and-Observability-Setup.docx'
document.save(output)
print(f'Created {output.name}: {len(document.paragraphs)} paragraphs, {len(document.tables)} tables')
